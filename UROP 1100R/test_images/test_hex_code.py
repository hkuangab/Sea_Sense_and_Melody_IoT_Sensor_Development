# 导入必要的库
from PIL import Image  # 用于图像处理
import numpy as np  # 用于数值计算
from collections import Counter  # 用于统计颜色出现频率
import cv2  # OpenCV库，用于图像显示
from typing import List, Dict  # 类型提示，提高代码可读性
import os  # 操作系统接口，用于文件操作
from datetime import datetime  # 日期时间处理


def lightweight_color_extraction(image_path: str, num_colors: int = 5) -> List[Dict]:
    """
    轻量级颜色提取函数（内存高效版本）
    
    参数:
        image_path (str): 输入图片的文件路径
        num_colors (int): 需要提取的颜色数量，默认为5种
    
    返回:
        List[Dict]: 包含颜色信息的字典列表，每个字典包含：
            - 'hex': HEX颜色代码（字符串，如'#FF0000'）
            - 'rgb': RGB颜色值（元组，如(255, 0, 0)）
            - 'percentage': 该颜色在图片中的占比（浮点数，百分比）
    """
    try:
        # 加载图片并转换为RGB模式
        img = Image.open(image_path)
        img = img.convert('RGB')
        
        # 缩小图片尺寸以节省内存
        img.thumbnail((200, 200))
        
        # 获取所有像素点的颜色数据
        pixels = list(img.getdata())
        
        # 简化颜色空间，减少颜色数量
        def simplify_color(rgb, step=32):
            r, g, b = rgb
            return (r//step*step, g//step*step, b//step*step)
        
        simplified = [simplify_color(p) for p in pixels]
        
        # 统计每种简化后颜色的出现次数
        color_counts = Counter(simplified)
        
        # 获取出现频率最高的num_colors种颜色
        most_common = color_counts.most_common(num_colors)
        
        # 将颜色信息转换为需要的格式
        colors = []
        for rgb, count in most_common:
            hex_color = '#{:02x}{:02x}{:02x}'.format(*rgb).upper()
            percentage = (count / len(pixels)) * 100
            colors.append({
                'hex': hex_color,
                'rgb': rgb,
                'percentage': round(percentage, 2)
            })
        
        return colors
        
    except Exception as e:
        print(f"Error: {e}")
        return []


def save_all_colors_to_one_excel(all_colors_data: Dict[str, List[Dict]], 
                                 output_file: str = "all_colors_combined.xlsx"):
    """
    将所有图片的颜色HEX代码保存到一个Excel文件中
    
    参数:
        all_colors_data (Dict[str, List[Dict]]): 字典，键为图片名称，值为颜色列表
        output_file (str): 输出的Excel文件名，默认为"all_colors_combined.xlsx"
    
    返回:
        bool: 保存成功返回True，失败返回False
    """
    try:
        # 尝试导入openpyxl库
        try:
            from openpyxl import Workbook
            from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        except ImportError:
            print("Error: openpyxl library is required")
            print("Please run: pip3 install openpyxl")
            return False
        
        # 创建工作簿
        wb = Workbook()
        
        # 创建汇总工作表
        summary_ws = wb.active
        summary_ws.title = "Summary"
        
        # 设置列宽
        summary_ws.column_dimensions['A'].width = 20
        summary_ws.column_dimensions['B'].width = 15
        summary_ws.column_dimensions['C'].width = 25
        
        # 创建标题样式
        title_font = Font(name='Arial', size=16, bold=True, color='FFFFFF')
        title_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        title_alignment = Alignment(horizontal='center', vertical='center')
        
        # 写入标题
        summary_ws.merge_cells('A1:C1')
        title_cell = summary_ws['A1']
        title_cell.value = "All Images Color HEX Codes"
        title_cell.font = title_font
        title_cell.fill = title_fill
        title_cell.alignment = title_alignment
        
        # 写入表头
        header_font = Font(name='Arial', size=12, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        
        headers = ["Image Name", "HEX Code", "Color Preview"]
        for col, header in enumerate(headers, 1):
            cell = summary_ws.cell(row=2, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # 写入汇总数据
        row = 3
        for image_name, colors in all_colors_data.items():
            for i, color_info in enumerate(colors):
                # 只在第一行显示图片名称
                if i == 0:
                    summary_ws.cell(row=row, column=1, value=image_name)
                
                # 写入HEX代码
                summary_ws.cell(row=row, column=2, value=color_info['hex'])
                
                # 写入颜色预览
                hex_code = color_info['hex'].lstrip('#')
                color_fill = PatternFill(start_color=hex_code, end_color=hex_code, fill_type='solid')
                preview_cell = summary_ws.cell(row=row, column=3, value="        ")
                preview_cell.fill = color_fill
                
                row += 1
            
            # 在不同图片之间添加空行
            row += 1
        
        # 为每个图片创建独立的工作表
        for image_name, colors in all_colors_data.items():
            # 创建工作表（工作表名称不能超过31个字符，不能包含特殊字符）
            ws_name = image_name[:31].replace(':', '_').replace('\\', '_').replace('/', '_').replace('*', '_').replace('?', '_').replace('[', '_').replace(']', '_')
            ws = wb.create_sheet(title=ws_name)
            
            # 设置列宽
            ws.column_dimensions['A'].width = 15
            ws.column_dimensions['B'].width = 25
            ws.column_dimensions['C'].width = 12
            ws.column_dimensions['D'].width = 25
            
            # 写入标题
            ws.merge_cells('A1:D1')
            title_cell = ws['A1']
            title_cell.value = f"Color HEX Codes for {image_name}"
            title_cell.font = title_font
            title_cell.fill = title_fill
            title_cell.alignment = title_alignment
            
            # 写入表头
            headers = ["No.", "HEX Code", "Percentage", "Color Preview"]
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=2, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            
            # 写入颜色数据
            for i, color_info in enumerate(colors, 1):
                # 序号
                ws.cell(row=i+2, column=1, value=i)
                
                # HEX代码
                ws.cell(row=i+2, column=2, value=color_info['hex'])
                
                # 百分比
                ws.cell(row=i+2, column=3, value=f"{color_info['percentage']}%")
                
                # 颜色预览
                hex_code = color_info['hex'].lstrip('#')
                color_fill = PatternFill(start_color=hex_code, end_color=hex_code, fill_type='solid')
                preview_cell = ws.cell(row=i+2, column=4, value="        ")
                preview_cell.fill = color_fill
        
        # 删除默认创建的空工作表
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
        
        # 保存Excel文件
        wb.save(output_file)
        
        print(f"✓ All color HEX codes saved to one Excel file: {output_file}")
        print(f"  File size: {os.path.getsize(output_file) / 1024:.1f} KB")
        
        return True
        
    except Exception as e:
        print(f"✗ Error saving all colors to one Excel file: {e}")
        return False


def save_hex_only_to_excel(colors: List[Dict], image_name: str, output_file: str = "hex_colors.xlsx"):
    """
    只保存单个图片的HEX颜色代码到Excel文件（简化版本）
    
    参数:
        colors (List[Dict]): 颜色信息列表
        image_name (str): 图片名称
        output_file (str): 输出的Excel文件名
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    except ImportError:
        print("Error: openpyxl library is required")
        return False
    
    wb = Workbook()
    ws = wb.active
    ws.title = image_name[:31]  # 工作表名称最多31个字符
    
    # 设置列宽
    ws.column_dimensions['A'].width = 15
    ws.column_dimensions['B'].width = 25
    
    # 创建标题
    title_font = Font(name='Arial', size=14, bold=True, color='FFFFFF')
    title_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    title_alignment = Alignment(horizontal='center', vertical='center')
    
    ws.merge_cells('A1:B1')
    title_cell = ws['A1']
    title_cell.value = f"HEX Color Codes for {image_name}"
    title_cell.font = title_font
    title_cell.fill = title_fill
    title_cell.alignment = title_alignment
    
    # 创建表头
    header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    
    headers = ["HEX Code", "Color Preview"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
    
    # 写入HEX代码和颜色预览
    for row, color_info in enumerate(colors, 3):
        hex_color = color_info['hex']
        
        # HEX代码
        hex_cell = ws.cell(row=row, column=1, value=hex_color)
        hex_cell.alignment = Alignment(horizontal='center')
        
        # 颜色预览
        hex_code = hex_color.lstrip('#')
        color_fill = PatternFill(start_color=hex_code, end_color=hex_code, fill_type='solid')
        preview_cell = ws.cell(row=row, column=2, value="        ")
        preview_cell.fill = color_fill
        preview_cell.alignment = Alignment(horizontal='center', vertical='center')
    
    # 保存文件
    wb.save(output_file)
    return True


def print_colors_in_terminal(colors: List[Dict], image_name: str = ""):
    """
    在终端中打印颜色信息
    
    参数:
        colors (List[Dict]): 颜色信息列表
        image_name (str): 图片名称
    """
    if image_name:
        print(f"\nColors for {image_name}:")
    
    print("="*60)
    for i, color_info in enumerate(colors, 1):
        hex_color = color_info['hex']
        rgb = color_info['rgb']
        percentage = color_info['percentage']
        print(f"{i:2d}. HEX: {hex_color:<8} RGB: {rgb} Percentage: {percentage:5.1f}%")


def check_openpyxl_installed():
    """
    检查openpyxl库是否已安装
    """
    try:
        import openpyxl
        return True
    except ImportError:
        return False


def main():
    """
    主函数：批量处理多张图片的颜色提取，并将所有HEX代码保存到一个Excel文件
    """
    # 检查openpyxl是否安装
    if not check_openpyxl_installed():
        print("Warning: openpyxl is not installed. Excel export will not work.")
        print("To install: pip3 install openpyxl")
        return
    
    # 定义图片列表
    image_files = ["hk", "ny", "sh", "paris"]
    image_dir = "/home/hkuangab/Desktop/UROP 1100R/test_images/"
    
    # 询问是否显示图片
    show_images = input("Do you want to display images? (y/n): ").lower().strip() == 'y'
    
    # 存储所有图片的颜色数据
    all_colors_data = {}
    
    # 遍历处理每张图片
    for i, image_file in enumerate(image_files, 1):
        # 构建完整的图片路径
        image_path = f"{image_dir}{image_file}.jpeg"
        
        print(f"\n{'='*60}")
        print(f"Processing image {i}/{len(image_files)}: {image_file}.jpeg")
        print(f"{'='*60}")
        
        # 检查图片文件是否存在
        if not os.path.exists(image_path):
            print(f"Error: Image file not found: {image_path}")
            print("Skipping to next image...")
            continue
        
        print(f"Image found: {image_path}")
        
        # 提取图片颜色
        print("Extracting colors from image...")
        colors = lightweight_color_extraction(image_path, num_colors=10)
        
        if not colors:
            print("No colors extracted. Skipping this image...")
            continue
        
        # 在终端显示提取的颜色
        print_colors_in_terminal(colors, image_file)
        
        # 保存到字典中
        all_colors_data[image_file] = colors
        
        # 可选显示原始图片
        if show_images:
            try:
                img = cv2.imread(image_path)
                if img is not None:
                    height, width = img.shape[:2]
                    max_display_size = 800
                    if max(height, width) > max_display_size:
                        scale = max_display_size / max(height, width)
                        new_width = int(width * scale)
                        new_height = int(height * scale)
                        img = cv2.resize(img, (new_width, new_height))
                    
                    cv2.imshow(f'Image: {image_file}', img)
                    print("Press any key to close the image window...")
                    cv2.waitKey(0)
                    cv2.destroyAllWindows()
            except Exception as e:
                print(f"Note: Could not display image. {e}")
        
        print(f"✓ Completed processing for {image_file}.jpeg")
    
    # 将所有颜色保存到一个Excel文件中
    if all_colors_data:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_file = f"all_colors_combined_{timestamp}.xlsx"
        
        if save_all_colors_to_one_excel(all_colors_data, output_file):
            print(f"\n✓ Successfully saved all color HEX codes to: {output_file}")
        else:
            print("\n✗ Failed to save to Excel file")
    else:
        print("\n✗ No colors extracted from any image.")
    
    # 完成提示
    print("\n" + "="*60)
    print("Batch color extraction completed!")
    print(f"Processed {len(all_colors_data)} images")
    print("="*60)


# 如果直接运行此脚本，执行主函数
if __name__ == "__main__":
    main()